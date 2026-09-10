"""Core deliverable rendering and structural verification.

The engine is deliberately domain-neutral. A domain pack supplies the
template declaration and section names; callers supply model prose and typed,
already-verified values. The engine owns the file bytes, structural checks,
hashes, and ledger records.
"""

from __future__ import annotations

import hashlib
import io
import re
import zipfile
from collections.abc import Callable, Mapping
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

import yaml
from docx import Document

from contracts import Clearance, LedgerEventEnvelope, Taint, build_event, idempotency_key, stable_id


MAX_TITLE_LENGTH = 255
MAX_SECTION_LENGTH = 200_000
MAX_VALUE_LENGTH = 4_096
MAX_ARTIFACT_BYTES = 100_000_000
_NAME_RE = re.compile(r"^[A-Za-z][A-Za-z0-9_.-]{0,127}$")
_IDENTITY_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_.-]{0,127}$")
_PLACEHOLDER_RE = re.compile(r"\{\{([A-Za-z][A-Za-z0-9_.-]{0,127})\}\}")
_CLEARANCE_RANK = {
    Clearance.public: 0,
    Clearance.internal: 1,
    Clearance.restricted: 2,
    Clearance.secret: 3,
}
_TAINT_RANK = {Taint.clean: 0, Taint.untrusted: 1, Taint.contaminated: 2}


class DeliverableError(RuntimeError):
    """A deliverable was rejected or could not be safely committed."""


class DeliverableLedger:
    @property
    def events(self) -> tuple[LedgerEventEnvelope, ...]: ...

    @property
    def head_hash(self) -> str | None: ...

    def append(self, event: LedgerEventEnvelope) -> Any: ...


@dataclass(frozen=True, slots=True)
class TemplateDefinition:
    template_id: str
    version: str
    file_format: str
    required_sections: tuple[str, ...]
    structural_check_required: bool
    visual_check_required: bool
    values_section: str | None = None
    allowed_value_names: tuple[str, ...] = ()
    required_value_names: tuple[str, ...] = ()


@dataclass(frozen=True, slots=True)
class DeterministicValue:
    """A named value that may be inserted into a deliverable.

    Values are supplied by a verified fact or deterministic computation. The
    model can refer to ``name`` in prose, but cannot supply the value itself.
    """

    name: str
    value_text: str
    unit: str | None
    value_origin: str
    source_refs: tuple[str, ...]
    evidence_refs: tuple[str, ...]
    verification_refs: tuple[str, ...]
    confidence: float
    clearance: Clearance
    taint: Taint
    derivation: dict[str, Any] | None
    verified: bool

    def validate(self) -> None:
        if not _NAME_RE.fullmatch(self.name):
            raise DeliverableError("deterministic value name is invalid")
        _bounded_text(self.value_text, MAX_VALUE_LENGTH, "deterministic value")
        if self.unit is not None:
            _bounded_text(self.unit, 128, "deterministic value unit")
        if self.value_origin not in {"verified_fact", "deterministic_calculation"}:
            raise DeliverableError("deterministic value origin is not permitted")
        _require_refs(self.source_refs, "deterministic value source")
        _require_refs(self.evidence_refs, "deterministic value evidence")
        _require_refs(self.verification_refs, "deterministic value verification")
        if type(self.confidence) not in (int, float) or not 0 <= self.confidence <= 1:
            raise DeliverableError("deterministic value confidence is invalid")
        if self.value_origin == "deterministic_calculation":
            if not isinstance(self.derivation, dict) or not self.derivation:
                raise DeliverableError("deterministic calculation requires derivation")
            inputs = self.derivation.get("input_refs")
            if not isinstance(inputs, (list, tuple)) or not inputs or not all(isinstance(item, str) and item.strip() for item in inputs):
                raise DeliverableError("deterministic calculation requires input references")
        if not self.verified:
            raise DeliverableError(f"deterministic value {self.name!r} is not verified")
        if self.taint == Taint.contaminated:
            raise DeliverableError(f"deterministic value {self.name!r} is contaminated")


@dataclass(frozen=True, slots=True)
class DeliverableRequest:
    task_id: str
    template_id: str
    title: str
    prose_sections: Mapping[str, str]
    values: tuple[DeterministicValue, ...]
    source_refs: tuple[str, ...]
    evidence_refs: tuple[str, ...]
    verification_refs: tuple[str, ...]
    clearance: Clearance
    taint: Taint
    confidence: float
    idempotency_key: str


@dataclass(frozen=True, slots=True)
class VisualCheckResult:
    status: str
    reason: str = ""


@dataclass(frozen=True, slots=True)
class DeliverableArtifact:
    artifact_id: str
    task_id: str
    template_id: str
    template_version: str
    title: str
    media_type: str
    file_format: str
    path: Path
    content_hash: str
    byte_size: int
    status: str
    verification_status: str
    structural_check: str
    visual_check: str
    source_refs: tuple[str, ...]
    evidence_refs: tuple[str, ...]
    verification_refs: tuple[str, ...]
    deterministic_value_refs: tuple[str, ...]
    confidence: float
    clearance: Clearance
    taint: Taint
    preview_blocks: tuple[tuple[str, str], ...]
    derivation: dict[str, Any]
    ledger_event_ref: str

    @property
    def approval_state(self) -> str:
        return "pending" if self.status in {"verified_draft", "needs_review"} else "unavailable"

    @property
    def approval_blocking_reasons(self) -> tuple[str, ...]:
        reasons: list[str] = []
        if self.verification_status != "passed":
            reasons.append("independent verification is not passed")
        if self.structural_check != "passed":
            reasons.append("structural artifact check is not passed")
        if self.visual_check not in {"passed", "not_required"}:
            reasons.append("visual artifact check is not passed")
        return tuple(reasons)


class LocalArtifactStore:
    """Bounded local artifact storage used by the first local Node slice."""

    def __init__(self, root: str | Path) -> None:
        self.root = Path(root).resolve()
        self.root.mkdir(parents=True, exist_ok=True)

    def write(self, artifact_id: str, file_format: str, content: bytes) -> Path:
        if not _IDENTITY_RE.fullmatch(artifact_id):
            raise DeliverableError("artifact identity is invalid")
        if file_format != "docx":
            raise DeliverableError("the local artifact store supports only DOCX in this slice")
        if len(content) > MAX_ARTIFACT_BYTES:
            raise DeliverableError("artifact exceeds the local size limit")
        target = self.root / f"{artifact_id}.docx"
        temporary = self.root / f".{artifact_id}.tmp"
        try:
            temporary.write_bytes(content)
            temporary.replace(target)
        except OSError as exc:
            raise DeliverableError("artifact could not be committed to local storage") from exc
        return target

    def read(self, artifact_id: str) -> bytes:
        if not _IDENTITY_RE.fullmatch(artifact_id):
            raise DeliverableError("artifact identity is invalid")
        target = self.root / f"{artifact_id}.docx"
        try:
            content = target.read_bytes()
        except OSError as exc:
            raise DeliverableError("artifact is not available in local storage") from exc
        if len(content) > MAX_ARTIFACT_BYTES:
            raise DeliverableError("stored artifact exceeds the local size limit")
        return content


class DeliverableEngine:
    """Render pack-declared deliverables from verified inputs."""

    def __init__(
        self,
        *,
        template_path: str | Path,
        artifact_store: LocalArtifactStore,
        ledger: DeliverableLedger,
        visual_checker: Callable[[Path], VisualCheckResult] | None = None,
        task_state_reader: Callable[[str], str] | None = None,
        actor_id: str = "deliverable-engine.local",
    ) -> None:
        self.template_path = Path(template_path).resolve()
        self.artifact_store = artifact_store
        self.ledger = ledger
        self.visual_checker = visual_checker
        self.task_state_reader = task_state_reader
        self.actor_id = actor_id
        self._records: dict[str, DeliverableArtifact] = {}
        self._templates = self._load_templates()

    def render(self, request: DeliverableRequest) -> DeliverableArtifact:
        template = self._templates.get(request.template_id)
        if template is None:
            raise DeliverableError("requested deliverable template is not declared by the domain pack")
        self._validate_request(request, template)
        if self.task_state_reader is not None:
            state = self.task_state_reader(request.task_id)
            if state not in {"rendering", "awaiting_review", "deliverable_verified"}:
                raise DeliverableError("deliverable rendering is not authorized in the current task state")

        artifact_id = stable_id("deliverable-artifact", request.task_id, request.template_id, request.idempotency_key)
        checked_key = idempotency_key("deliverable.artifact.checked", artifact_id)
        existing_checked = self._event_by_key(checked_key)
        if existing_checked is not None:
            existing = self._records.get(artifact_id)
            if existing is None:
                raise DeliverableError("the sealed artifact exists but its local record is unavailable")
            return existing

        rendered_sections = {
            section: _replace_named_values(request.prose_sections[section], {value.name: value.value_text for value in request.values})
            for section in template.required_sections
        }
        content = self._render_docx(request.title, rendered_sections, request.values, template)
        structural_status, structural_reason = _structural_check(content, template, rendered_sections, request.values)
        if structural_status != "passed":
            raise DeliverableError(structural_reason)

        artifact_id = stable_id("deliverable-artifact", request.task_id, request.template_id, request.idempotency_key)
        path = self.artifact_store.write(artifact_id, template.file_format, content)
        visual = VisualCheckResult("not_required")
        if template.visual_check_required:
            visual = self.visual_checker(path) if self.visual_checker is not None else VisualCheckResult("unavailable", "No local visual renderer is configured.")
            if visual.status not in {"passed", "unavailable", "failed"}:
                raise DeliverableError("visual checker returned an invalid status")

        status = "verified_draft" if visual.status in {"passed", "not_required"} else "needs_review"
        verification_status = "passed" if visual.status in {"passed", "not_required"} else "needs_review"
        content_hash = hashlib.sha256(content).hexdigest()
        source_refs = _ordered_unique((*request.source_refs, *(ref for value in request.values for ref in value.source_refs)))
        evidence_refs = _ordered_unique((*request.evidence_refs, *(ref for value in request.values for ref in value.evidence_refs)))
        verification_refs = _ordered_unique((*request.verification_refs, *(ref for value in request.values for ref in value.verification_refs)))
        taint = max((request.taint, *(value.taint for value in request.values)), key=lambda item: _TAINT_RANK[item])
        confidence = min((request.confidence, *(value.confidence for value in request.values)))
        derivation = {value.name: value.derivation for value in request.values if value.derivation is not None}
        preview_blocks = [{"kind": "title", "text": request.title}]
        preview_blocks.extend({"kind": "section", "text": text} for text in rendered_sections.values())
        approval_blocking_reasons: list[str] = []
        if verification_status != "passed":
            approval_blocking_reasons.append("independent verification is not passed")
        if structural_status != "passed":
            approval_blocking_reasons.append("structural artifact check is not passed")
        if visual.status not in {"passed", "not_required"}:
            approval_blocking_reasons.append("visual artifact check is not passed")
        provenance = {
            "source_refs": list(source_refs),
            "evidence_refs": list(evidence_refs),
            "verification_refs": list(verification_refs),
            "deterministic_value_refs": [value.name for value in request.values],
            "confidence": confidence,
            "clearance": request.clearance.value,
            "taint": taint.value,
            "derivation": derivation,
        }
        common_payload = {
            "artifact_id": artifact_id,
            "task_id": request.task_id,
            "template_id": template.template_id,
            "template_version": template.version,
            "title": request.title,
            "media_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "format": template.file_format,
            "content_hash": content_hash,
            "byte_size": len(content),
            "source_refs": list(source_refs),
            "evidence_refs": list(evidence_refs),
            "verification_refs": list(verification_refs),
            "deterministic_value_refs": [value.name for value in request.values],
            "provenance": provenance,
            "confidence": confidence,
            "clearance": request.clearance.value,
            "taint": taint.value,
            "derivation": derivation,
            "preview_blocks": preview_blocks,
        }
        staged = self._append(
            event_type="artifact.staged",
            task_id=request.task_id,
            payload={**common_payload, "status": "staged"},
            contract="DeliverableArtifact",
            key=idempotency_key("deliverable.artifact.staged", artifact_id),
            clearance=request.clearance,
        )
        try:
            checked = self._append(
                event_type="artifact.checked",
                task_id=request.task_id,
                payload={
                    **common_payload,
                    "status": status,
                    "verification_status": verification_status,
                    "checks": {
                        "structural": structural_status,
                        "visual": visual.status,
                        "visual_reason": visual.reason,
                    },
                    "approval_state": "pending",
                    "approval_blocking_reasons": approval_blocking_reasons,
                },
                contract="DeliverableArtifactCheck",
                key=checked_key,
                clearance=request.clearance,
            )
        except Exception as exc:
            # The staged file remains quarantined behind a ledger record. The
            # gateway exposes only artifacts with a committed checked event.
            raise DeliverableError("artifact check was not committed to the ledger") from exc

        artifact = DeliverableArtifact(
            artifact_id=artifact_id,
            task_id=request.task_id,
            template_id=template.template_id,
            template_version=template.version,
            title=request.title,
            media_type=common_payload["media_type"],
            file_format=template.file_format,
            path=path,
            content_hash=content_hash,
            byte_size=len(content),
            status=status,
            verification_status=verification_status,
            structural_check=structural_status,
            visual_check=visual.status,
            source_refs=source_refs,
            evidence_refs=evidence_refs,
            verification_refs=verification_refs,
            deterministic_value_refs=tuple(value.name for value in request.values),
            confidence=confidence,
            clearance=request.clearance,
            taint=taint,
            preview_blocks=tuple((item["kind"], item["text"]) for item in preview_blocks),
            derivation=derivation,
            ledger_event_ref=checked.event_id,
        )
        self._records[artifact_id] = artifact
        return artifact

    def get(self, artifact_id: str) -> DeliverableArtifact:
        try:
            return self._records[artifact_id]
        except KeyError as exc:
            raise DeliverableError("artifact record is unavailable") from exc

    def _load_templates(self) -> dict[str, TemplateDefinition]:
        try:
            payload = yaml.safe_load(self.template_path.read_text(encoding="utf-8"))
        except (OSError, UnicodeError, yaml.YAMLError) as exc:
            raise DeliverableError("domain-pack deliverable template declaration could not be loaded") from exc
        raw_templates = payload.get("templates") if isinstance(payload, dict) else None
        if not isinstance(raw_templates, list):
            raise DeliverableError("domain-pack deliverable templates must be a list")
        result: dict[str, TemplateDefinition] = {}
        for raw in raw_templates:
            if not isinstance(raw, dict):
                raise DeliverableError("domain-pack deliverable template entry is invalid")
            template_id = raw.get("id")
            version = raw.get("version", "0.1")
            file_format = raw.get("format")
            sections = raw.get("required_sections")
            if not isinstance(template_id, str) or not _NAME_RE.fullmatch(template_id):
                raise DeliverableError("domain-pack deliverable template ID is invalid")
            if not isinstance(version, str) or not version.strip():
                raise DeliverableError("domain-pack deliverable template version is invalid")
            if file_format != "docx" or not isinstance(sections, list) or not sections:
                raise DeliverableError("the local deliverable slice requires a DOCX template with sections")
            names = tuple(section for section in sections if isinstance(section, str) and _NAME_RE.fullmatch(section))
            if len(names) != len(sections) or len(set(names)) != len(names):
                raise DeliverableError("domain-pack deliverable sections are invalid")
            if template_id in result:
                raise DeliverableError("domain-pack deliverable template IDs must be unique")
            value_bindings = raw.get("value_bindings", [])
            allowed: list[str] = []
            required: list[str] = []
            if not isinstance(value_bindings, list):
                raise DeliverableError("domain-pack value bindings must be a list")
            for binding in value_bindings:
                if not isinstance(binding, dict) or not isinstance(binding.get("name"), str) or not _NAME_RE.fullmatch(binding["name"]):
                    raise DeliverableError("domain-pack value binding is invalid")
                allowed.append(binding["name"])
                if binding.get("required") is True:
                    required.append(binding["name"])
            result[template_id] = TemplateDefinition(
                template_id=template_id,
                version=version,
                file_format=file_format,
                required_sections=names,
                structural_check_required=raw.get("structural_check") in {True, "required"},
                visual_check_required=raw.get("visual_check") in {True, "required"},
                values_section=raw.get("values_section") if isinstance(raw.get("values_section"), str) else None,
                allowed_value_names=tuple(allowed),
                required_value_names=tuple(required),
            )
        return result

    def _validate_request(self, request: DeliverableRequest, template: TemplateDefinition) -> None:
        # Task IDs are repository identities and may be UUIDs generated by
        # the orchestrator, including UUIDs whose first character is a digit.
        # Keep the stricter name grammar for template and value names, but do
        # not reject a valid orchestrator task at the deliverable boundary.
        if not _IDENTITY_RE.fullmatch(request.task_id):
            raise DeliverableError("task identity is invalid")
        if request.template_id != template.template_id:
            raise DeliverableError("template identity mismatch")
        _bounded_text(request.title, MAX_TITLE_LENGTH, "deliverable title")
        _require_refs(request.source_refs, "deliverable source")
        _require_refs(request.evidence_refs, "deliverable evidence")
        _require_refs(request.verification_refs, "deliverable verification")
        if type(request.confidence) not in (int, float) or not 0 <= request.confidence <= 1:
            raise DeliverableError("deliverable confidence is invalid")
        if request.taint == Taint.contaminated:
            raise DeliverableError("contaminated deliverable inputs cannot be rendered")
        if not isinstance(request.idempotency_key, str) or not request.idempotency_key.strip() or len(request.idempotency_key) > 256:
            raise DeliverableError("deliverable idempotency key is invalid")
        if set(request.prose_sections) != set(template.required_sections):
            missing = set(template.required_sections) - set(request.prose_sections)
            unknown = set(request.prose_sections) - set(template.required_sections)
            raise DeliverableError(f"deliverable sections do not match the pack declaration: missing={sorted(missing)}, unknown={sorted(unknown)}")
        for section in template.required_sections:
            _bounded_text(request.prose_sections[section], MAX_SECTION_LENGTH, f"deliverable section {section}")
            if not request.prose_sections[section].strip():
                raise DeliverableError(f"deliverable section {section} is empty")
        names = [value.name for value in request.values]
        if len(set(names)) != len(names):
            raise DeliverableError("deterministic value names must be unique")
        if template.allowed_value_names and any(name not in template.allowed_value_names for name in names):
            raise DeliverableError("deterministic value is not declared by the pack template")
        if any(name not in names for name in template.required_value_names):
            raise DeliverableError("a required deterministic value is missing")
        for value in request.values:
            value.validate()
            if _CLEARANCE_RANK[value.clearance] > _CLEARANCE_RANK[request.clearance]:
                raise DeliverableError("deliverable would drop value clearance")
            if _TAINT_RANK[value.taint] > _TAINT_RANK[request.taint]:
                raise DeliverableError("deliverable would drop value taint")
        placeholders = {
            name
            for section in request.prose_sections.values()
            for name in _PLACEHOLDER_RE.findall(section)
        }
        if placeholders - set(names):
            raise DeliverableError(f"unbound named value: {sorted(placeholders - set(names))[0]}")

    @staticmethod
    def _render_docx(
        title: str,
        sections: Mapping[str, str],
        values: tuple[DeterministicValue, ...],
        template: TemplateDefinition,
    ) -> bytes:
        document = Document()
        document.core_properties.title = title
        document.add_heading(title, level=0)
        for section in template.required_sections:
            document.add_heading(_section_label(section), level=1)
            document.add_paragraph(sections[section])
            if template.values_section == section and values:
                table = document.add_table(rows=1, cols=3)
                table.style = "Table Grid"
                headers = table.rows[0].cells
                headers[0].text = "Name"
                headers[1].text = "Value"
                headers[2].text = "Unit"
                for value in values:
                    cells = table.add_row().cells
                    cells[0].text = value.name
                    cells[1].text = value.value_text
                    cells[2].text = value.unit or ""
        output = io.BytesIO()
        document.save(output)
        return output.getvalue()

    def _append(
        self,
        *,
        event_type: str,
        task_id: str,
        payload: dict[str, Any],
        contract: str,
        key: str,
        clearance: Clearance,
    ) -> LedgerEventEnvelope:
        existing = self._event_by_key(key)
        if existing is not None:
            return existing
        event = build_event(
            event_type=event_type,
            task_id=task_id,
            actor_id=self.actor_id,
            actor_type="service",
            payload_contract=contract,
            payload_version="1.0",
            payload=payload,
            clearance=clearance,
            idempotency=key,
            sequence=len(self.ledger.events),
            previous_event_hash=self.ledger.head_hash,
        )
        try:
            committed = self.ledger.append(event)
        except Exception as exc:
            raise DeliverableError("deliverable ledger append failed") from exc
        if isinstance(committed, LedgerEventEnvelope):
            return committed
        return self._event_by_key(key) or event

    def _event_by_key(self, key: str) -> LedgerEventEnvelope | None:
        return next((event for event in self.ledger.events if event.idempotency_key == key), None)


def _bounded_text(value: Any, maximum: int, label: str) -> str:
    if not isinstance(value, str) or "\x00" in value or len(value) > maximum:
        raise DeliverableError(f"{label} is invalid")
    if any(ord(char) < 32 and char not in "\n\r\t" for char in value):
        raise DeliverableError(f"{label} contains unsupported control characters")
    return value


def _require_refs(values: tuple[str, ...], label: str) -> None:
    if not values or any(not isinstance(value, str) or not value.strip() or len(value) > 512 for value in values):
        raise DeliverableError(f"{label} references are incomplete")


def _ordered_unique(values: tuple[str, ...]) -> tuple[str, ...]:
    return tuple(dict.fromkeys(value for value in values if isinstance(value, str) and value.strip()))


def _replace_named_values(text: str, values: Mapping[str, str]) -> str:
    def replace(match: re.Match[str]) -> str:
        name = match.group(1)
        if name not in values:
            raise DeliverableError(f"unbound named value: {name}")
        return values[name]

    return _PLACEHOLDER_RE.sub(replace, text)


def _section_label(section: str) -> str:
    return section.replace("_", " ").strip().title()


def _structural_check(
    content: bytes,
    template: TemplateDefinition,
    sections: Mapping[str, str],
    values: tuple[DeterministicValue, ...],
) -> tuple[str, str]:
    if not template.structural_check_required:
        return "not_required", ""
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as package:
            names = set(package.namelist())
            required = {"[Content_Types].xml", "word/document.xml", "word/styles.xml"}
            if not required.issubset(names):
                return "failed", "DOCX package is missing required OOXML parts"
            if any(name.endswith(("vbaProject.bin", ".exe", ".dll")) for name in names):
                return "failed", "executable or macro content is not permitted in the first-scope artifact"
            for name in names:
                if not name.endswith(".rels"):
                    continue
                root = ElementTree.fromstring(package.read(name))
                for relationship in root:
                    if relationship.attrib.get("TargetMode") == "External" or relationship.attrib.get("Target", "").lower().startswith(("http:", "https:")):
                        return "failed", "external document relationships are not permitted"
            document_root = ElementTree.fromstring(package.read("word/document.xml"))
            text = "".join(document_root.itertext())
            if "{{" in text or "}}" in text:
                return "failed", "unresolved named value placeholder remains in the DOCX"
            for section, rendered in sections.items():
                if _section_label(section) not in text and rendered not in text:
                    return "failed", f"required section {section!r} was not found in the DOCX"
            for value in values:
                if value.value_text not in text:
                    return "failed", f"deterministic value {value.name!r} was not found in the DOCX"
    except (zipfile.BadZipFile, KeyError, UnicodeError, ElementTree.ParseError) as exc:
        return "failed", "DOCX structural check could not safely read the package"
    return "passed", ""


__all__ = [
    "DeliverableArtifact",
    "DeliverableEngine",
    "DeliverableError",
    "DeliverableRequest",
    "DeterministicValue",
    "LocalArtifactStore",
    "TemplateDefinition",
    "VisualCheckResult",
]
