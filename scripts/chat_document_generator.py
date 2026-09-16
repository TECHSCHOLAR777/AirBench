"""Standalone freeform document generator for the desktop Chat feature.

Reads one JSON object from stdin:
    {
      "format": "docx" | "xlsx" | "pptx" | "pdf",
      "title": str,
      "sections": [{"heading": str, "body": str}, ...],
      "table": {"headers": [str, ...], "rows": [[str, ...], ...]} | null,
      "output_path": str
    }

Writes the rendered document to output_path. Exits 0 on success, or prints
an error message to stderr and exits 1.

This is intentionally separate from src/airbench/delivery/engine.py, which
renders deliverables bound to task records and fixed required-section
templates. This script renders arbitrary, chat-requested content instead.
"""
from __future__ import annotations

import json
import sys
import textwrap


def _read_payload() -> dict:
    raw = sys.stdin.read()
    return json.loads(raw)


def _write_docx(payload: dict, output_path: str) -> None:
    from docx import Document

    document = Document()
    document.core_properties.title = payload["title"]
    document.add_heading(payload["title"], level=0)
    for section in payload.get("sections", []):
        document.add_heading(section.get("heading", ""), level=1)
        document.add_paragraph(section.get("body", ""))
    table = payload.get("table")
    if table and table.get("headers"):
        grid = document.add_table(rows=1, cols=len(table["headers"]))
        grid.style = "Table Grid"
        for cell, header in zip(grid.rows[0].cells, table["headers"]):
            cell.text = str(header)
        for row in table.get("rows", []):
            cells = grid.add_row().cells
            for cell, value in zip(cells, row):
                cell.text = str(value)
    document.save(output_path)


def _write_xlsx(payload: dict, output_path: str) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Document"
    row = 1
    sheet.cell(row=row, column=1, value=payload["title"])
    row += 1
    for section in payload.get("sections", []):
        row += 1
        sheet.cell(row=row, column=1, value=section.get("heading", ""))
        row += 1
        sheet.cell(row=row, column=1, value=section.get("body", ""))
    table = payload.get("table")
    if table and table.get("headers"):
        row += 2
        for column, header in enumerate(table["headers"], start=1):
            sheet.cell(row=row, column=column, value=str(header))
        for values in table.get("rows", []):
            row += 1
            for column, value in enumerate(values, start=1):
                sheet.cell(row=row, column=column, value=str(value))
    workbook.save(output_path)


def _write_pptx(payload: dict, output_path: str) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_slide.shapes.title.text = payload["title"]

    body_layout = presentation.slide_layouts[1]
    for section in payload.get("sections", []):
        slide = presentation.slides.add_slide(body_layout)
        slide.shapes.title.text = section.get("heading", "")
        if len(slide.placeholders) > 1:
            slide.placeholders[1].text_frame.text = section.get("body", "")

    table = payload.get("table")
    if table and table.get("headers"):
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        rows = len(table.get("rows", [])) + 1
        columns = len(table["headers"])
        grid = slide.shapes.add_table(
            rows, columns, Inches(0.5), Inches(1.5), Inches(9.0), Inches(0.4 * rows)
        ).table
        for column, header in enumerate(table["headers"]):
            grid.cell(0, column).text = str(header)
        for row_index, values in enumerate(table.get("rows", []), start=1):
            for column, value in enumerate(values):
                grid.cell(row_index, column).text = str(value)

    presentation.save(output_path)


def _pdf_escape(value: str) -> str:
    safe = value.encode("latin-1", "replace").decode("latin-1")
    return safe.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)").replace("\r", " ").replace("\n", " ")


def _pdf_page_content(lines: list[tuple[str, int]], page_number: int, page_count: int) -> bytes:
    commands = ["BT", "/F1 18 Tf", "54 742 Td"]
    first = True
    for text, size in lines:
        if not first:
            commands.append("0 -15 Td")
        commands.append(f"/F1 {size} Tf")
        commands.append(f"({_pdf_escape(text)}) Tj")
        first = False
    commands.extend(["/F1 8 Tf", "1 0 0 1 54 36 Tm", f"(Page {page_number} of {page_count}) Tj", "ET"])
    return "\n".join(commands).encode("latin-1", "replace")


def _write_pdf(payload: dict, output_path: str) -> None:
    """Dependency-free Helvetica-only PDF writer, generalized from the
    approach in src/airbench/delivery/engine.py::_render_pdf for arbitrary
    (rather than template-bound) titles/sections/tables."""
    lines: list[tuple[str, int]] = [(payload["title"], 18)]
    for section in payload.get("sections", []):
        lines.append((section.get("heading", ""), 14))
        for paragraph in section.get("body", "").splitlines() or [""]:
            wrapped = textwrap.wrap(paragraph, width=96, break_long_words=True, break_on_hyphens=False) or [""]
            lines.extend((line, 10) for line in wrapped)
    table = payload.get("table")
    if table and table.get("headers"):
        lines.append((" | ".join(str(header) for header in table["headers"]), 10))
        for row in table.get("rows", []):
            lines.append((" | ".join(str(value) for value in row), 10))

    page_lines: list[list[tuple[str, int]]] = []
    current: list[tuple[str, int]] = []
    remaining = 42
    for line in lines:
        cost = 2 if line[1] >= 14 else 1
        if current and remaining < cost:
            page_lines.append(current)
            current = []
            remaining = 42
        current.append(line)
        remaining -= cost
    if current:
        page_lines.append(current)
    if not page_lines:
        page_lines = [[("(empty document)", 10)]]

    objects: list[bytes] = [b"", b""]
    page_object_ids: list[int] = []
    content_object_ids: list[int] = []
    for _ in page_lines:
        page_object_ids.append(len(objects) + 1)
        objects.append(b"")
        content_object_ids.append(len(objects) + 1)
        objects.append(b"")
    font_object_id = len(objects) + 1
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    objects[0] = b"<< /Type /Catalog /Pages 2 0 R >>"
    kids = " ".join(f"{object_id} 0 R" for object_id in page_object_ids)
    objects[1] = f"<< /Type /Pages /Kids [{kids}] /Count {len(page_object_ids)} >>".encode("ascii")
    for index, page in enumerate(page_lines):
        content = _pdf_page_content(page, index + 1, len(page_lines))
        objects[page_object_ids[index] - 1] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font_object_id} 0 R >> >> "
            f"/Contents {content_object_ids[index]} 0 R >>"
        ).encode("ascii")
        objects[content_object_ids[index] - 1] = (
            f"<< /Length {len(content)} >>\nstream\n".encode("ascii") + content + b"\nendstream"
        )

    output = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for object_id, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{object_id} 0 obj\n".encode("ascii"))
        output.extend(body)
        output.extend(b"\nendobj\n")
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode("ascii"))
    output.extend(b"".join(f"{offset:010d} 00000 n \n".encode("ascii") for offset in offsets[1:]))
    output.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode("ascii")
    )
    with open(output_path, "wb") as handle:
        handle.write(bytes(output))


_WRITERS = {
    "docx": _write_docx,
    "xlsx": _write_xlsx,
    "pptx": _write_pptx,
    "pdf": _write_pdf,
}


def main() -> int:
    payload = _read_payload()
    file_format = payload.get("format")
    writer = _WRITERS.get(file_format)
    if writer is None:
        print(f"Unsupported document format: {file_format!r}", file=sys.stderr)
        return 1
    output_path = payload["output_path"]
    try:
        writer(payload, output_path)
    except ImportError as error:
        print(f"Missing Python dependency for {file_format}: {error}", file=sys.stderr)
        return 1
    except Exception as error:  # noqa: BLE001 - surface any renderer failure to the caller
        print(f"Failed to render {file_format} document: {error}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
